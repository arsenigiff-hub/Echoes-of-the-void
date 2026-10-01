from pathlib import Path
import re
import json
import hashlib
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

ROOT = Path(r'D:\godot\GodotGame-main')
WORK = ROOT / '.codex_qa/work_report'
PATH = ROOT / 'Отчет о проделанной работе ECHOES OF THE VOID.docx'
d = Document(PATH)
figs = {
1: ('Основная игровая сцена с контрастным освещением', ['10_corridor_overview.png'], 'Обзорный кадр основной сцены с отдельной диагностической камеры. Видны организация прохода, геометрия интерьера и зоны голубого и красного освещения.'),
2: ('Обратный ракурс интерьера корабля', ['11_corridor_reverse.png'], 'Второй обзорный кадр той же сцены с диагностической камеры, направленной к противоположной части помещения. Ресурсы и освещение уровня сохранены без изменения.'),
3: ('Главное меню с моделью Земли', ['01_menu.png'], 'Фактический кадр main_menu.tscn до запуска новой игры. На экране присутствуют модель Земли и четыре кнопки меню.'),
4: ('Фактическое изображение вступительной сцены', ['03_intro.png'], 'Кадр intro_cutscene.tscn из сохраненного положения камеры. Камера направлена на близкую геометрию интерьера; виден экранный дизеринг. Ошибка имени анимации, описанная в разделе, не исправлялась при съемке.'),
5: ('Вид от первого лица после приземления персонажа', ['09_main_after_landing.png'], 'Кадр камеры Player после расчетного приземления на опорную поверхность. Камера направлена вдоль помещения, фонарь выключен; изображение показывает стационарное освещение.'),
6: ('Режим работы с радиоприемником', ['07_radio_focus.png'], 'Крупный план после завершения Tween к CameraTarget. В этом режиме скорость игрока равна нулю, а обычное перемещение заблокировано.'),
7: ('Радиоприемник после команды поворота ручки', ['08_radio_knob.png'], 'Кадр того же прибора после команды D, вызывающей поворот Radio_TuningKnob_001 на 60 градусов вокруг локальной оси. Числовое изменение ориентации подтверждено проверкой состояния узла; частота и захват сигнала пока не реализованы.'),
8: ('Сравнение одного ракурса с фонарем и без него', ['14_flashlight_off.png', '15_flashlight_on.png'], 'Верхний кадр — фонарь выключен; нижний — включен. Положение камеры и стационарные источники одинаковы. Меняется только видимость SpotLight3D персонажа.'),
9: ('Сравнение изображения меню с дизерингом и без него', ['12_menu_dither_on.png', '13_menu_dither_off.png'], 'Верхний кадр — ColorRect с шейдером включен, dither_size = 4 и color_depth = 12; нижний — этот экранный слой скрыт. Положение камеры и угол Земли одинаковы, чтобы показать действие эффекта.'),
}

before_code = [p.text for p in d.paragraphs if p.style.name == 'Code']
inserted = []
for p in list(d.paragraphs):
    m = re.fullmatch(r'Место для скриншота (\d+)', p.text)
    if not m:
        continue
    num = int(m.group(1))
    title, files, explanation = figs[num]
    cap = Paragraph(p._p.getnext(), p._parent)
    guide = Paragraph(cap._p.getnext(), p._parent)
    p.clear()
    p.style = d.styles['Normal']
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.space_before = Pt(6)
    pf.space_after = Pt(4)
    pf.line_spacing = 1
    pf.keep_with_next = True
    pf.keep_together = True
    for shd in p._p.xpath('./w:pPr/w:shd'):
        shd.getparent().remove(shd)
    for idx, name in enumerate(files):
        src = WORK/'screens'/name
        if not src.exists():
            raise FileNotFoundError(src)
        if idx:
            p.add_run().add_break()
        r = p.add_run()
        shape = r.add_picture(str(src), width=Cm(11.5 if len(files)>1 else 15))
        desc = shape._inline.docPr
        desc.set('descr', f'Рисунок {num}. {title}. Кадр {idx+1}.')
        inserted.append({'figure': num, 'file': name, 'sha256': hashlib.sha256(src.read_bytes()).hexdigest()})
    cap.text = f'Рисунок {num} — {title}'
    cap.paragraph_format.keep_with_next = True
    guide.text = explanation
    guide.paragraph_format.keep_with_next = False

rename = {'Приложение В Порядок подготовки скриншотов': 'Приложение В Перечень скриншотов'}
for p in d.paragraphs:
    if p.text in rename:
        p.text = rename[p.text]
    elif p.text.startswith('В основной части предусмотрены девять мест'):
        p.text = 'В основной части размещены девять пронумерованных иллюстраций, содержащих одиннадцать кадров проекта. Скриншоты получены непосредственно из области отображения Godot при загрузке существующих сцен. Для обзорных видов использована диагностическая камера, для игровых действий — камера персонажа. В сравнительных кадрах менялась только видимость фонаря или экранного слоя шейдера.'
    elif p.text.startswith('При подготовке кадра меню следует'):
        p.text = 'Главное меню снималось отдельным запуском main_menu.tscn, поскольку главной сценой проекта остается main.tscn. Изображение вступления показывает текущее положение камеры и существующий экранный эффект, включая ограничение незавершенной анимации. Крупные планы радио сняты после завершения перехода камеры к прибору.'
    elif p.text.startswith('Кадр после поворота ручки должен'):
        p.text = 'Пары кадров фонаря и дизеринга сняты с одинаковым ракурсом внутри каждой пары. При сравнении шейдера вращение Земли временно остановлено в диагностическом экземпляре, чтобы различие относилось к экранному эффекту. Исходные игровые сцены и скрипты при подготовке иллюстраций не изменялись.'

last_table = d.tables[-1]
rows = [
('1', 'Обзор основной сцены с контрастным светом', 'Раздел 3'),
('2', 'Обратный ракурс интерьера', 'Раздел 4'),
('3', 'Главное меню с Землей и кнопками', 'Раздел 5'),
('4', 'Текущий кадр вступления с дизерингом', 'Раздел 6'),
('5', 'Вид камеры персонажа после приземления', 'Раздел 7'),
('6', 'Крупный план радио после перехода камеры', 'Раздел 9'),
('7', 'Тот же прибор после команды D', 'Раздел 9'),
('8', 'Два кадра с выключенным и включенным фонарем', 'Раздел 10'),
('9', 'Два кадра меню с шейдером и без него', 'Раздел 11'),
]
for row, values in zip(last_table.rows[1:], rows):
    for cell, value in zip(row.cells, values):
        p = cell.paragraphs[0]
        p.clear()
        r = p.add_run(value)
        r.font.size = Pt(10)

assert len(inserted) == 11, len(inserted)
assert before_code == [p.text for p in d.paragraphs if p.style.name == 'Code']
assert not any(p.text.startswith('Место для скриншота') for p in d.paragraphs)
d.save(PATH)
(WORK/'screenshots_inserted.json').write_text(json.dumps(inserted,ensure_ascii=False,indent=2),encoding='utf-8')
print('Figures 9; embedded frames', len(d.inline_shapes), '; code paragraphs preserved', len(before_code))
