from pathlib import Path
import re
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:\godot\GodotGame-main')
WORK = ROOT / '.codex_qa/work_report'
OUTPUT = ROOT / 'Отчет о проделанной работе ECHOES OF THE VOID.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(2), Cm(2)
s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
s.footer_distance = Cm(1)

def font(style, name, size, bold=False):
    st = d.styles[style]
    st.font.name = name
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    st.element.rPr.rFonts.set(qn('w:cs'), name)
    return st

normal = font('Normal', 'Times New Roman', 12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.first_line_indent = Cm(0.7)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.widow_control = True
for name, size in [('Title', 22), ('Heading 1', 16), ('Heading 2', 13)]:
    st = font(name, 'Times New Roman', size, True)
    st.paragraph_format.first_line_indent = Cm(0)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True
for border in d.styles['Title'].element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
caption = font('Caption', 'Times New Roman', 11)
caption.paragraph_format.first_line_indent = Cm(0)
caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_after = Pt(5)
for name in ['Code', 'ScreenSlot', 'ScreenGuide', 'ContentsLine']:
    if name not in d.styles:
        d.styles.add_style(name, 1)
code = font('Code', 'Consolas', 8)
code.paragraph_format.first_line_indent = Cm(0)
code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
code.paragraph_format.line_spacing = 1
code.paragraph_format.space_after = Pt(0)
code.paragraph_format.space_before = Pt(0)
code.paragraph_format.keep_with_next = False
code.paragraph_format.keep_together = True
slot = font('ScreenSlot', 'Times New Roman', 11)
slot.paragraph_format.first_line_indent = Cm(0)
slot.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
slot.paragraph_format.space_before = Pt(18)
slot.paragraph_format.space_after = Pt(45)
slot.paragraph_format.keep_with_next = True
guide = font('ScreenGuide', 'Times New Roman', 10)
guide.paragraph_format.first_line_indent = Cm(0)
guide.paragraph_format.space_after = Pt(8)
guide.paragraph_format.keep_together = True
cont = font('ContentsLine', 'Times New Roman', 11)
cont.paragraph_format.first_line_indent = Cm(0)
cont.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
cont.paragraph_format.line_spacing = 1
cont.paragraph_format.space_after = Pt(3)

p = s.footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Cm(0)
run = p.add_run()
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
run._r.addnext(field)

def shade(p, color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    p._p.get_or_add_pPr().append(shd)

def table(lines):
    rows = [[x.strip() for x in ln.strip().strip('|').split('|')] for ln in lines]
    rows = [r for r in rows if not all(re.match(r'^:?-+:?$', c) for c in r)]
    t = d.add_table(rows=0, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    n = len(rows[0])
    widths = {2:[6,11],3:[5,4,8],4:[5,3,4,5],5:[5.4,2.9,4.1,2.3,2.3]}.get(n,[17/n]*n)
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge = OxmlElement('w:' + side)
        for k, v in [('val','single'),('sz','4'),('color','D9D9D9')]:
            edge.set(qn('w:' + k),v)
        borders.append(edge)
    t._tbl.tblPr.append(borders)
    for ri, vals in enumerate(rows):
        row = t.add_row()
        no_split=OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        if ri == 0:
            repeat=OxmlElement('w:tblHeader')
            row._tr.get_or_add_trPr().append(repeat)
        for ci, (c, value) in enumerate(zip(row.cells, vals)):
            c.width = Cm(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'),'D9E2F3' if ri == 0 else ('F5F7FA' if ri % 2 == 0 else 'FFFFFF'))
            pr.append(shd)
            mar = OxmlElement('w:tcMar')
            for k in ['top','bottom','left','right']:
                e = OxmlElement('w:' + k)
                e.set(qn('w:w'),'95')
                e.set(qn('w:type'),'dxa')
                mar.append(e)
            pr.append(mar)
            p=c.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(value)
            r.font.size=Pt(10)
            r.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)

listing_stats=[]
def listing(rel):
    src=(ROOT/rel).read_text(encoding='utf-8').splitlines()
    listing_stats.append({'file':rel,'lines':len(src)})
    for idx,line in enumerate(src,1):
        p=d.add_paragraph(style='Code')
        p.add_run(f'{idx:03d}  {line.expandtabs(4)}')
    d.add_paragraph().paragraph_format.space_after=Pt(0)

lines=(WORK/'report.md').read_text(encoding='utf-8').splitlines()
i=0
in_contents=False
slots=0
while i<len(lines):
    line=lines[i]
    if not line.strip():
        i+=1; continue
    if line.startswith('# '):
        d.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):
        text=line[3:]
        if text=='Содержание' or text.startswith('1 Общая') or text.startswith('Приложение '):
            d.add_page_break()
        d.add_paragraph(text,style='Heading 1')
        in_contents=text=='Содержание'
    elif line.startswith('### '):
        d.add_paragraph(line[4:],style='Heading 2')
    elif line.startswith('[[LISTING '):
        listing(line[10:-2])
    elif line.startswith('[[SCREEN '):
        number,title,instruction=line[9:-2].split('|',2)
        p=d.add_paragraph(f'Место для скриншота {number}',style='ScreenSlot')
        shade(p,'F2F2F2')
        p=d.add_paragraph(f'Рисунок {number} — {title}',style='Caption')
        p.paragraph_format.keep_with_next=True
        d.add_paragraph(instruction,style='ScreenGuide')
        slots+=1
    elif line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].startswith('|'):
            block.append(lines[i]);i+=1
        table(block)
        continue
    elif line.startswith('```'):
        i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            d.add_paragraph(lines[i],style='Code');i+=1
    else:
        p=d.add_paragraph(line,style='ContentsLine' if in_contents else 'Normal')
        if line.startswith('Подробное описание'):
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
            p.runs[0].font.size=Pt(14)
        elif line.startswith('Дата отчета'):
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.space_after=Pt(18)
    i+=1

d.core_properties.title='Отчет о разработке игры ECHOES OF THE VOID'
d.core_properties.subject='Выполненная работа, описание реализации, полные листинги и результаты проверки'
d.core_properties.author=''
d.core_properties.keywords='Godot, GDScript, ECHOES OF THE VOID, отчет'
d.save(OUTPUT)
(WORK/'build_stats.json').write_text(json.dumps({'output':str(OUTPUT),'screenshots_embedded':0,'screenshot_slots':slots,'listings':listing_stats,'paragraphs':len(d.paragraphs),'tables':len(d.tables)},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUTPUT)
print('paragraphs',len(d.paragraphs),'tables',len(d.tables),'slots',slots,'listing lines',sum(x['lines'] for x in listing_stats))
